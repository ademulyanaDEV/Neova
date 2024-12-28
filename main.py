import os
import random
import string
from fpdf import FPDF
import csv
from PIL import Image, ImageDraw

def generate_random_content(size):
    return ''.join(random.choices(string.ascii_letters + string.digits, k=size))

def create_random_file(folder_name, file_type, file_size, file_index):
    random_name = ''.join(random.choices(string.ascii_letters + string.digits, k=10))  # Generate random name
    file_name = f"{file_index}_{random_name}.{file_type if file_type != 'nft' else 'png'}"
    file_path = os.path.join(folder_name, file_name)

    if file_type == "txt":
        with open(file_path, "w") as file:
            while file.tell() < file_size:
                file.write(generate_random_content(1024))  # Write in 1KB chunks
    elif file_type == "csv":
        with open(file_path, "w", newline="") as file:
            writer = csv.writer(file)
            while os.path.getsize(file_path) < file_size:
                writer.writerow([generate_random_content(10) for _ in range(5)])
    elif file_type == "pdf":
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("Arial", size=12)
        content_size = 0
        while content_size < file_size:
            content = generate_random_content(100)
            pdf.cell(200, 10, txt=content, ln=True)
            content_size = len(pdf.output(dest='S').encode('latin1'))
        pdf.output(file_path)
    elif file_type == "docx":
        try:
            from docx import Document
            doc = Document()
            temp_path = f"{file_path}.tmp"  # Temporary file path for size calculation
            doc.save(temp_path)  # Save an initial empty document
            
            while os.path.getsize(temp_path) < file_size:
                doc.add_paragraph(generate_random_content(100))
                doc.save(temp_path)
            
            os.rename(temp_path, file_path)  # Rename temp file to final file name
        except ImportError:
            print("Please install python-docx library to generate Word files.")
    elif file_type == "nft":
        image = Image.new('RGB', (500, 500), (255, 255, 255))
        draw = ImageDraw.Draw(image)
        for _ in range(100):
            x1, y1 = random.randint(0, 499), random.randint(0, 499)
            x2, y2 = random.randint(0, 499), random.randint(0, 499)
            x1, x2 = sorted([x1, x2])  # Ensure x1 <= x2
            y1, y2 = sorted([y1, y2])  # Ensure y1 <= y2
            color = tuple(random.choices(range(256), k=3))
            draw.rectangle([x1, y1, x2, y2], fill=color, outline=color)
        image.save(file_path)

    print(f"File created: {file_path}")

def main():
    folder_name = "TEMPLATE_FILE"
    os.makedirs(folder_name, exist_ok=True)

    num_files = int(input("Enter the number of files to create: "))
    print("Select the file type:")
    print("1. TXT")
    print("2. CSV")
    print("3. PDF")
    print("4. DOCX")
    print("5. NFT (Image)")
    file_type_option = input("Enter your choice (1/2/3/4/5): ")
    file_type_map = {"1": "txt", "2": "csv", "3": "pdf", "4": "docx", "5": "nft"}
    file_type = file_type_map.get(file_type_option, "txt")
    file_size_mb = 1
    file_size = file_size_mb * 1024 * 1024  # Convert MB to bytes

    for i in range(1, num_files + 1):
        create_random_file(folder_name, file_type, file_size, i)

if __name__ == "__main__":
    main()
